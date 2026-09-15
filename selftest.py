# -*- coding: utf-8 -*-
"""自测：生成样例文档，验证解析与三条规则判定（不需要飞书凭证和网络）"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document

from fsaudit.docparse import parse_doc
from fsaudit.rules import (audit_text, load_ai_terms, load_negative_terms,
                           split_keywords, col_letter)

FAILURES = []


def check(name, cond, detail=""):
    tag = "PASS" if cond else "FAIL"
    print("  [%s] %s %s" % (tag, name, ("- " + detail) if detail and not cond else ""))
    if not cond:
        FAILURES.append(name)


def make_docx(path, title, body):
    d = Document()
    d.add_paragraph(title)
    for p in body:
        d.add_paragraph(p)
    d.save(path)


def main():
    tmp = tempfile.mkdtemp()
    terms = load_ai_terms()

    print("== 解析 + 规则自测 ==")

    # 1. 应通过：我方品牌先出现，无AI痕迹、无违禁词
    p1 = os.path.join(tmp, "pass.docx")
    make_docx(p1, "蓝海新品发布", ["蓝海新品今天正式发布。", "性能全面领先。", "友商竞品X也在跟进。"])
    title, text = parse_doc(p1)
    res = audit_text(text, split_keywords("蓝海, 蓝海科技"), terms, ["AI生成", "免责声明"], ["竞品X"])
    check("规则全过样例", res["passed"], str(res["reasons"]))

    # 2. 规则一不通过：出现痕迹话术（模型名已不拦截，2026-09-13 口径收敛）
    p2 = os.path.join(tmp, "r1.docx")
    make_docx(p2, "蓝海体验", ["蓝海新品体验很好。", "本文以下由AI生成。", "据说 DeepSeek 也赞不绝口。"])
    _, text = parse_doc(p2)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("规则一命中痕迹话术", (not res["passed"]) and any("规则一" in r and "以下由ai生成" in r.lower() for r in res["reasons"]), str(res["reasons"]))

    # 3. 规则二不通过：出现「免责声明」和「AI生成」
    p3 = os.path.join(tmp, "r2.docx")
    make_docx(p3, "蓝海介绍", ["蓝海公司介绍。", "本AI生成内容仅供参考。", "免责声明：本文不构成建议。"])
    _, text = parse_doc(p3)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    hits = [r for r in res["reasons"] if "规则二" in r]
    check("规则二双命中", (not res["passed"]) and len(hits) == 1 and "AI生成" in hits[0] and "免责声明" in hits[0], str(res["reasons"]))

    # 3b. 规则二空白变体：CSDN 水印「（注：部分内容可能由 AI 生成）」夹空格也要命中（2026-09-13 真实漏判案例）
    p3b = os.path.join(tmp, "r2b.docx")
    make_docx(p3b, "蓝海介绍", ["蓝海公司介绍。", "（注：部分内容可能由 AI 生成）"])
    _, text = parse_doc(p3b)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("规则二空格变体命中", (not res["passed"]) and any("规则二" in r and "AI生成" in r for r in res["reasons"]), str(res["reasons"]))

    # 3c. 规则二开头署名/时间戳：「作者：xx」「时间：xx年x月」不通过（2026-09-13 新口径）
    p3c = os.path.join(tmp, "r2c.docx")
    make_docx(p3c, "蓝海体验报告", ["作者：张三", "时间：2026年9月", "蓝海新品体验很好。"])
    _, text = parse_doc(p3c)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("规则二开头署名时间戳", (not res["passed"]) and any("规则二" in r and "署名" in r and "作者" in r for r in res["reasons"]), str(res["reasons"]))

    # 3d. 规则二开头署名防误伤：正文提及「作者」、散文「时间：」不拦截
    p3d = os.path.join(tmp, "r2d.docx")
    make_docx(p3d, "蓝海体验报告", ["时间：是最好的证明。", "本文作者认为蓝海不错。"])
    _, text = parse_doc(p3d)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("规则二署名防误伤", res["passed"], str(res["reasons"]))

    # 4. 规则三不通过：竞品先出现
    p4 = os.path.join(tmp, "r3.docx")
    make_docx(p4, "竞品X对比评测", ["竞品X最近很火。", "对比来看蓝海更胜一筹。"])
    _, text = parse_doc(p4)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"])
    check("规则三竞品先现", (not res["passed"]) and any("规则三" in r for r in res["reasons"]), str(res["reasons"]))

    # 5. 规则三：我方关键词未出现
    p5 = os.path.join(tmp, "r3b.docx")
    make_docx(p5, "随便一篇文章", ["这是一篇没有提到品牌的文章。"])
    _, text = parse_doc(p5)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("规则三我方未出现", (not res["passed"]) and any("未在文中出现" in r for r in res["reasons"]), str(res["reasons"]))

    # 6. 多关键词：任一我方词先于竞品即可
    p6 = os.path.join(tmp, "r3c.docx")
    make_docx(p6, "文章", ["蓝海科技发布了新品。", "竞品X随后跟进。"])
    _, text = parse_doc(p6)
    res = audit_text(text, split_keywords("蓝海、蓝海科技"), terms, ["AI生成", "免责声明"], ["竞品X"])
    check("多关键词先现通过", res["passed"], str(res["reasons"]))

    # 7. 品牌名包含关系不算竞品
    p7 = os.path.join(tmp, "r3d.docx")
    make_docx(p7, "文章", ["蓝海集团旗下产品上市。"])
    _, text = parse_doc(p7)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [])
    check("包含关系归我方", res["passed"], str(res["reasons"]))

    # 8. 关键词拆分
    check("关键词拆分", split_keywords("蓝海, 蓝海科技、XYZ｜abc") == ["蓝海", "蓝海科技", "XYZ", "abc"])

    # 9. 列号转换
    check("列号转换", col_letter(1) == "A" and col_letter(26) == "Z" and col_letter(27) == "AA" and col_letter(53) == "BA")

    # ---- 规则八 · 竞品联系方式（2026-09-15 新增） ----
    neg = load_negative_terms()

    p8a = os.path.join(tmp, "r8a.docx")
    make_docx(p8a, "蓝海对比", ["蓝海新品正式发布，抢先体验。", "想深入了解竞品X，可拨打客服电话13812345678咨询。"])
    _, text = parse_doc(p8a)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"],
                     None, None, neg)
    check("规则八竞品手机号", (not res["passed"]) and any("规则八" in r and "手机号" in r for r in res["reasons"]), str(res["reasons"]))

    p8b = os.path.join(tmp, "r8b.docx")
    make_docx(p8b, "蓝海对比", ["蓝海新品正式发布。", "竞品X详情可加微信ABC12345获取。"])
    _, text = parse_doc(p8b)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"],
                     None, None, neg)
    check("规则八竞品微信号", (not res["passed"]) and any("规则八" in r and "微信号" in r for r in res["reasons"]), str(res["reasons"]))

    p8c = os.path.join(tmp, "r8c.docx")
    make_docx(p8c, "蓝海发布", ["蓝海新品正式发布。", "如需选购可咨询我方客服400-123-4567，或在官网 www.lanhai.com 下单。"])
    _, text = parse_doc(p8c)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [],
                     None, None, neg)
    check("规则八我方联系方式不拦", res["passed"], str(res["reasons"]))

    p8d = os.path.join(tmp, "r8d.docx")
    make_docx(p8d, "蓝海介绍", ["蓝海新品正式发布，我方客服电话400-123-4567。",
                                "行业动态方面，据多家媒体报道，今年上半年整个行业呈现明显复苏态势，"
                                "多家品牌陆续发布新品，市场关注度持续升温，展会现场人气火爆，热闹非凡。",
                                "在这样的大背景下，同期举办的某大型展会上，也有竞品X的新品亮相，获得不少关注。"])
    _, text = parse_doc(p8d)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"],
                     None, None, neg)
    check("规则八距离外不误伤", res["passed"], str(res["reasons"]))

    # ---- 规则九 · 品牌负面描述本地快速路径（2026-09-15 新增） ----
    p9a = os.path.join(tmp, "r9a.docx")
    make_docx(p9a, "蓝海对比", ["蓝海新品表现稳健。", "相比之下竞品X近期翻车不断，口碑下滑。"])
    _, text = parse_doc(p9a)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"],
                     None, None, neg)
    check("规则九竞品负面词", (not res["passed"]) and any("规则九" in r and "竞品X" in r for r in res["reasons"]), str(res["reasons"]))

    p9b = os.path.join(tmp, "r9b.docx")
    make_docx(p9b, "蓝海争议", ["蓝海新品引发热议。", "但也有网友吐槽蓝海定价偏高，直呼智商税。"])
    _, text = parse_doc(p9b)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [],
                     None, None, neg)
    check("规则九我方负面词", (not res["passed"]) and any("规则九" in r and "蓝海" in r for r in res["reasons"]), str(res["reasons"]))

    p9c = os.path.join(tmp, "r9c.docx")
    make_docx(p9c, "蓝海观察", ["蓝海新品上市，性能、口碑、销量全面开花，成为今年的现象级产品，市场反响十分热烈。",
                                "另一方面，行业内个别小作坊此前曾出现翻车传闻，监管已介入处理。"])
    _, text = parse_doc(p9c)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], [],
                     None, None, neg)
    check("规则九窗口外不误伤", res["passed"], str(res["reasons"]))

    p9d = os.path.join(tmp, "r9d.docx")
    make_docx(p9d, "蓝海对比", ["蓝海新品正式发布，续航表现优于竞品X，性价比更高。"])
    _, text = parse_doc(p9d)
    res = audit_text(text, ["蓝海"], terms, ["AI生成", "免责声明"], ["竞品X"],
                     None, None, neg)
    check("规则九客观对比不误伤", res["passed"], str(res["reasons"]))

    print("")
    if FAILURES:
        print("自测未通过 %d 项：%s" % (len(FAILURES), "、".join(FAILURES)))
        sys.exit(1)
    print("自测全部通过 ✔")


if __name__ == "__main__":
    main()
